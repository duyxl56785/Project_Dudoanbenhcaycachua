from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "report" / "bao_cao_nhom_tomato_template.docx"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def set_cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, color="FFFFFF")
        shade(cell, "1F4E79")
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    doc.add_paragraph()
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
for name, size, color in (("Title", 24, "1F4E79"), ("Heading 1", 16, "1F4E79"), ("Heading 2", 12, "2F5597")):
    styles[name].font.name = "Aptos Display"
    styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Nhận diện bệnh qua ảnh lá cà chua")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Báo cáo bài tập lớn học phần Học máy và Khai phá dữ liệu")
run.bold = True
run.font.size = Pt(12)
subtitle.add_run("\nTrường Đại học CMC | Nhóm 4 | Bản khung NV5")

doc.add_paragraph()
intro = doc.add_paragraph()
intro.add_run("Mục đích. ").bold = True
intro.add_run("Tài liệu này là khung báo cáo để nhóm điền kết quả đã truy vết từ code, checkpoint và Test split chung. Các trường đặt trong ngoặc vuông là chỗ cần cập nhật sau khi NV2, NV3 và NV4 bàn giao artifact.")

doc.add_heading("Mở đầu", level=1)
doc.add_paragraph("Đề tài phân loại ảnh lá cà chua vào bốn lớp Healthy, Early Blight, Late Blight và Bacterial Spot. Nhóm so sánh một CNN huấn luyện từ đầu với ResNet18 dùng Transfer Learning, đồng thời cung cấp demo Streamlit.")
doc.add_paragraph("Phạm vi của báo cáo là phân loại ảnh trong các lớp đã học. Kết quả trên PlantVillage không đại diện trực tiếp cho ảnh ngoài thực địa và chưa đủ cơ sở để kết luận khả năng phát hiện bệnh mới chớm.")

doc.add_heading("Phần I Tổng quan", level=1)
doc.add_heading("1.1 Bài toán", level=2)
doc.add_paragraph("[Mô tả nhu cầu, đầu vào, đầu ra và ý nghĩa của bốn lớp theo docs/DATASET_README.md.] ")
doc.add_heading("1.2 Mục tiêu", level=2)
for item in ["Chuẩn hóa pipeline dữ liệu và mapping nhãn dùng chung.", "Huấn luyện và đánh giá CNN cùng ResNet18 trên cùng Test split.", "Đối chiếu kết quả với demo và nêu rõ giới hạn triển khai."]:
    add_bullet(doc, item)

doc.add_heading("Phần II Phương pháp thực hiện", level=1)
doc.add_heading("2.1 Dữ liệu và preprocessing", level=2)
doc.add_paragraph("[Nguồn dữ liệu, số ảnh từng lớp, cách chia train/validation/test, kiểm tra trùng và augmentation. Điền từ bàn giao NV1.]")
doc.add_heading("2.2 CNN tự xây dựng", level=2)
doc.add_paragraph("[Kiến trúc, loss, optimizer, seed, số epoch, tiêu chí chọn checkpoint. Điền từ bàn giao NV2.]")
doc.add_heading("2.3 ResNet18 và Transfer Learning", level=2)
doc.add_paragraph("[Weights, feature extraction, fine-tuning, preprocessing và tiêu chí chọn checkpoint. Điền từ bàn giao NV3.]")
doc.add_heading("2.4 Demo", level=2)
doc.add_paragraph("[Cách chạy Streamlit, preprocessing theo từng model và ảnh minh họa. Điền từ bàn giao NV4.]")

doc.add_heading("Phần III Thực nghiệm và đánh giá", level=1)
doc.add_heading("3.1 Giao thức đánh giá", level=2)
doc.add_paragraph("NV5 dùng scripts/evaluate_models.py để kiểm tra hai prediction CSV có cùng sample ID và ground truth. Test chỉ được dùng sau khi checkpoint đã được chọn bằng Validation.")
add_table(doc, ["Mô hình", "Split", "Số mẫu", "Accuracy", "Macro F1", "Nguồn"], [["CNN", "Test", "[điền]", "[điền]", "[điền]", "results/cnn/"], ["ResNet18", "Test", "[điền]", "[điền]", "[điền]", "results/resnet18/"]])
doc.add_heading("3.2 Bảng so sánh", level=2)
doc.add_paragraph("[Chèn bảng results/comparison.csv sau khi chạy đánh giá thật. Không dùng số liệu minh họa.]")
doc.add_heading("3.3 Ma trận nhầm lẫn và phân tích lỗi", level=2)
doc.add_paragraph("[Chèn hai confusion matrix và các ảnh Test đúng/sai. Nêu cặp lớp nhầm nhiều nhất dựa trên results/analysis.md.]")
doc.add_heading("3.4 Điều kiện đo", level=2)
add_table(doc, ["Trường", "Giá trị"], [["Thiết bị", "[CPU/GPU]"], ["Python/PyTorch", "[điền phiên bản]"], ["Batch size", "[điền]"], ["Thời gian train/inference", "[điền từ metadata]"]])

doc.add_heading("Phần IV Kết luận và giới hạn", level=1)
doc.add_paragraph("[Kết luận chỉ dựa trên bảng so sánh thật. Nêu mô hình có kết quả phù hợp hơn trong Test split, các giới hạn của dữ liệu và hướng thu thập ảnh thực địa.] ")
for item in ["PlantVillage có điều kiện chụp được kiểm soát.", "Mô hình chỉ phân loại bốn lớp và có thể ép ảnh ngoài miền vào một lớp.", "Chưa có thí nghiệm độc lập về phát hiện bệnh mới chớm."]:
    add_bullet(doc, item)

doc.add_heading("Tài liệu tham khảo", level=1)
doc.add_paragraph("[Điền nguồn PlantVillage, PyTorch/torchvision, ResNet18 và các tài liệu nhóm đã sử dụng. Mỗi nguồn cần URL hoặc DOI kiểm tra được.]")
doc.add_heading("Phụ lục", level=1)
doc.add_paragraph("[Lệnh chạy, cấu hình, checksum hoặc commit của checkpoint, bộ câu hỏi phản biện và phân công thuyết trình.]")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Bản khung NV5 – cập nhật sau Gate 3 và Gate 5").font.size = Pt(9)
doc.save(OUT)
print(OUT)
