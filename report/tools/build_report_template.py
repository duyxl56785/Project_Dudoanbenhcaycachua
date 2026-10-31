from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "report" / "deliverables" / "bao_cao_nhom_tomato_two_models_nv5.docx"
ASSETS = ROOT / "report" / "assets"
CNN_MATRIX = ASSETS / "results" / "cnn_confusion_matrix.png"
RESNET_MATRIX = ASSETS / "results" / "resnet18_confusion_matrix.png"
CNN_DEMO = ASSETS / "nv4_demo" / "cnn_prediction_real.jpg"
RESNET_DEMO = ASSETS / "nv4_demo" / "resnet18_real_20261030.jpg"
COMPARE_DEMO = ASSETS / "nv4_demo" / "cnn_resnet18_compare_20261030.jpg"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def set_cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_text(cell, header, bold=True, color="FFFFFF")
        shade(cell, "1F4E79")
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    header_props.append(header_marker)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    doc.add_paragraph()
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)


def add_figure(doc, image_path, caption, alt_text, width=6.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    inline_shape = p.add_run().add_picture(str(image_path), width=Inches(width))
    inline_shape._inline.docPr.set("descr", alt_text)
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_p.add_run(caption)
    caption_run.italic = True
    caption_run.font.size = Pt(9)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
styles["Normal"].font.size = Pt(10.5)
for name, size, color in (("Title", 24, "000000"), ("Heading 1", 16, "000000"), ("Heading 2", 12, "000000")):
    styles[name].font.name = "Aptos Display"
    styles[name]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    styles[name]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor.from_string(color)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Nhận diện bệnh qua ảnh lá cà chua")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Báo cáo bài tập lớn học phần Học máy và Khai phá dữ liệu")
run.bold = True
run.font.size = Pt(12)
subtitle.add_run("\nTrường Đại học CMC | Nhóm 4 | Bản dự thảo cập nhật bởi NV5")

doc.add_paragraph()
intro = doc.add_paragraph()
intro.add_run("Mục đích. ").bold = True
intro.add_run("NV5 đã đối chiếu CNN và ResNet18 từ prediction CSV cùng Test split, kiểm hash checkpoint ResNet18 và ghi nguồn metadata. Kết quả chỉ áp dụng cho artifact PlantVillage hiện có.")

doc.add_heading("Mở đầu", level=1)
doc.add_paragraph("Đề tài phân loại ảnh lá cà chua vào bốn lớp Healthy, Early Blight, Late Blight và Bacterial Spot. Nhóm đánh giá CNN học từ đầu và ResNet18 Transfer Learning trên cùng Test split, đồng thời cung cấp demo Streamlit.")
doc.add_paragraph("Phạm vi là phân loại ảnh trong bốn lớp đã học. Kết quả PlantVillage không đại diện trực tiếp cho ảnh ngoài thực địa và chưa đủ cơ sở để kết luận khả năng phát hiện bệnh mới chớm.")

doc.add_heading("Phần I Tổng quan", level=1)
doc.add_heading("1.1 Bài toán", level=2)
doc.add_paragraph("Đầu vào là ảnh RGB của lá cà chua, đầu ra là một trong bốn nhãn theo configs/class_mapping.json. Đây là bài toán phân loại ảnh nhiều lớp, không phải phát hiện đối tượng hay phân đoạn ảnh.")
doc.add_heading("1.2 Mục tiêu", level=2)
for item in ["Chuẩn hóa pipeline dữ liệu và mapping nhãn dùng chung.", "Huấn luyện và đánh giá CNN cùng ResNet18 trên cùng Test split.", "Đối chiếu kết quả với demo và nêu rõ giới hạn triển khai."]:
    add_bullet(doc, item)

doc.add_heading("Phần II Phương pháp thực hiện", level=1)
doc.add_heading("2.1 Dữ liệu và preprocessing", level=2)
doc.add_paragraph("Bộ dữ liệu PlantVillage trong pipeline chung có 6.502 ảnh: 4.115 Train, 1.031 Validation và 1.356 Test. Train có thể augmentation; Validation và Test dùng tiền xử lý tất định. Mapping nhãn 0–3 là Healthy, Early Blight, Late Blight, Bacterial Spot. Nhóm cần xác nhận phiên bản manifest, seed và bằng chứng loại trùng trước khi nộp bản cuối.")
doc.add_heading("2.2 CNN tự xây dựng", level=2)
doc.add_paragraph("NV2 huấn luyện CNN từ đầu cho bốn lớp. Artifact PR #9 mô tả bốn khối Conv–BatchNorm–ReLU–MaxPool với số kênh 32 đến 256, Global Average Pooling, Dropout 0,3 và lớp Linear 4 đầu ra. Cấu hình được chọn theo Validation là E1 với learning rate 3e-4; checkpoint epoch 24 có Validation macro-F1 0,976836. NV5 chỉ dùng artifact đã chốt này để kiểm định Test.")
doc.add_heading("2.3 ResNet18 và Transfer Learning", level=2)
doc.add_paragraph("NV3 bàn giao checkpoint ResNet18 fine-tuned, metadata, prediction CSV Test, metrics, log và curve qua ZIP trong PR #7. NV5 kiểm SHA-256 checkpoint 61d44ba8ac02ee57c44cbd3835aaac2c657c97924735fea743a60df5f5fdf9e2. Stage fine-tuning, epoch 9, được chọn bằng Validation macro-F1 0,980859; ResNet18 dùng ImageNet pretraining và fine-tuning layer4.")
doc.add_heading("2.4 Demo", level=2)
doc.add_paragraph("NV4 đã kiểm chứng demo Streamlit với cả checkpoint CNN và ResNet18 thật: strict-load CPU, demo.check_consistency --model both PASS với max |delta| = 0 và metric ResNet18 Test tái hiện trên CPU. Ảnh chụp demo có nguồn NV4 được đưa vào phần 3.5; video trình diễn chưa được bàn giao.")

doc.add_heading("Phần III Thực nghiệm và đánh giá", level=1)
doc.add_heading("3.1 Giao thức đánh giá", level=2)
doc.add_paragraph("NV5 dùng scripts/evaluate_models.py để kiểm tra hai prediction CSV có cùng sample ID và ground truth. Test chỉ được dùng sau khi checkpoint đã được chọn bằng Validation.")
add_table(doc, ["Mô hình", "Split", "Số mẫu", "Accuracy", "Macro F1", "Nguồn"], [["CNN", "Test", "1.356", "0,969764", "0,965034", "results/cnn_metrics_nv2_handoff.json; PR #9"], ["ResNet18", "Test", "1.356", "0,986726", "0,983978", "results/resnet18_metrics_nv3_handoff.json; PR #7"]])
doc.add_heading("3.2 Bảng so sánh", level=2)
doc.add_paragraph("NV5 chạy scripts/evaluate_models.py với hai prediction CSV thật. Hai model có cùng 1.356 sample ID và ground truth sau khi chuẩn hóa tiền tố path Colab/Windows. ResNet18 cao hơn CNN 0,016962 Accuracy và 0,018944 macro-F1; đây là mô tả artifact, không phải chọn lại model trên Test.")
add_table(doc, ["Lớp", "F1 CNN", "F1 ResNet18", "Support Test"], [["Healthy", "0,990964", "0,993958", "330"], ["Early Blight", "0,927400", "0,962264", "215"], ["Late Blight", "0,952128", "0,985430", "375"], ["Bacterial Spot", "0,989643", "0,994259", "436"]])
doc.add_heading("3.3 Ma trận nhầm lẫn và phân tích lỗi", level=2)
doc.add_paragraph("CNN đúng 1.315/1.356 ảnh, ResNet18 đúng 1.338/1.356 ảnh. CNN nhầm nhiều nhất Early Blight → Late Blight (14) và Late Blight → Early Blight (12). ResNet18 giảm lỗi xuống 18 ảnh; lỗi nhiều nhất vẫn Early Blight → Late Blight (6). ImageNet pretraining và fine-tuning là diễn giải hợp lý cho kết quả cao hơn, nhưng cần nhiều seed và Test ngoài miền để khẳng định.")
add_figure(doc, CNN_MATRIX, "Hình 1. Ma trận nhầm lẫn CNN. Nguồn: prediction CSV NV2, tổng hợp bởi NV5.", "Ma trận nhầm lẫn CNN trên 1.356 ảnh Test.", width=5.9)
add_figure(doc, RESNET_MATRIX, "Hình 2. Ma trận nhầm lẫn ResNet18. Nguồn: prediction CSV NV3, tổng hợp bởi NV5.", "Ma trận nhầm lẫn ResNet18 trên 1.356 ảnh Test.", width=5.9)
doc.add_heading("3.4 Điều kiện đo", level=2)
add_table(doc, ["Trường", "Giá trị"], [["CNN", "389.924 tham số; train 1.118,7 giây; RTX 3050"], ["ResNet18", "11.178.564 tham số; 8.395.780 trainable; Tesla T4"], ["Checkpoint ResNet18", "fine-tuning epoch 9; SHA-256 đã kiểm"], ["Python/PyTorch", "CNN 3.12.10/2.14.1; ResNet18 3.13.16/2.11.0"], ["Thời gian ResNet18", "Không có artifact, để trống"]])
doc.add_heading("3.5 Minh chứng demo Streamlit", level=2)
doc.add_paragraph("NV4 đã strict-load hai checkpoint trên CPU và chạy demo.check_consistency --model both với max |delta| = 0. Các hình dưới đây là ảnh chụp giao diện trên một ảnh Test Bacterial Spot. Chúng minh họa cách sử dụng giao diện, không phải phép đo Accuracy toàn tập.")
add_figure(
    doc,
    COMPARE_DEMO,
    "Hình 3. Chế độ so sánh CNN và ResNet18. Nguồn: NV4 commit 2406b12, demo/screenshots/cnn_resnet18_compare_20261030.jpg.",
    "Ảnh chụp giao diện Streamlit so sánh CNN và ResNet18 trên một ảnh Test Bacterial Spot.",
)
add_figure(doc, CNN_DEMO, "Hình 4. CNN dự đoán Bacterial Spot 96,95%. Nguồn: NV4 commit 2406b12.", "Ảnh chụp giao diện Streamlit cho lượt dự đoán CNN Bacterial Spot 96,95%.", width=4.2)
add_figure(doc, RESNET_DEMO, "Hình 5. ResNet18 dự đoán Bacterial Spot 67,11%. Nguồn: NV4 commit 2406b12.", "Ảnh chụp giao diện Streamlit cho lượt dự đoán ResNet18 Bacterial Spot 67,11%.", width=4.2)
doc.add_paragraph("NV4 chưa bàn giao ảnh chụp riêng cho một dự đoán sai có nhãn thật hiển thị. Báo cáo vì vậy không gắn nhãn lại ảnh demo; các lỗi được trình bày bằng confusion matrix và prediction CSV.")

doc.add_heading("Phần IV Kết luận và giới hạn", level=1)
doc.add_paragraph("ResNet18 đạt macro-F1 0,983978, cao hơn CNN 0,965034 trên cùng 1.356 ảnh Test. Lỗi vẫn tập trung ở Early Blight và Late Blight. Cả hai checkpoint đã chọn bằng Validation; kết luận chỉ áp dụng PlantVillage và bản này chưa phải báo cáo nộp cuối nếu chưa có mẫu CMC chính thức, ảnh/video demo được nhóm xác nhận.")
for item in ["PlantVillage có điều kiện chụp được kiểm soát.", "Mô hình chỉ phân loại bốn lớp và có thể ép ảnh ngoài miền vào một lớp.", "Chưa có thí nghiệm độc lập về phát hiện bệnh mới chớm."]:
    add_bullet(doc, item)

doc.add_heading("Tài liệu tham khảo", level=1)
for reference in [
    "Hughes, D. P., & Salathé, M. (2015). An open access repository of images on plant health to enable the development of mobile disease diagnostics. arXiv:1511.08060. https://arxiv.org/abs/1511.08060",
    "He, K., Zhang, X., Ren, S., & Sun, J. (2016). Deep Residual Learning for Image Recognition. CVPR. DOI: 10.1109/CVPR.2016.90.",
    "scikit-learn developers. precision_recall_fscore_support documentation. https://scikit-learn.org/stable/modules/generated/sklearn.metrics.precision_recall_fscore_support.html",
    "Tài liệu nội bộ: docs/PHAN_CONG.md, docs/CNN_README.md, Issue #6, PR #7, PR #9, PR #11 và report/RESULTS_TRACEABILITY.md.",
]:
    doc.add_paragraph(reference)
doc.add_heading("Phụ lục", level=1)
doc.add_paragraph("Lệnh tái tạo bảng: xem docs/NV5_EVALUATION.md, dùng prediction CSV CNN/ResNet18 đã bàn giao cùng hai JSON handoff. Cập nhật comparison.csv, analysis.md, Word và slide từ cùng một đầu ra. Bộ 20 câu hỏi phản biện nằm tại docs/NV5_DEFENSE_QA.md.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Bản dự thảo NV5 – hai mô hình đã kiểm chứng trên Test chung").font.size = Pt(9)
doc.save(OUT)
print(OUT)
